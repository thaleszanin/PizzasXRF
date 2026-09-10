# -*- coding: utf-8 -*-
"""Leitura da planilha de concentrações.

É a planilha que o programa de concentrações exporta: um arquivo .xlsx
cuja aba **"Resultados"** traz uma amostra por linha e, por elemento, um
par de colunas (a concentração e o erro dela). O cabeçalho ocupa três
linhas:

    ----------+-------------+-------------+------------+---
              | Ca (mg/kg)  |             | Cl (mg/kg) |     <- 1: elemento
      LD      | -           |             | -          |     <- 2: limite
              | C           | Erro        | C          |     <- 3: o que é
    ----------+-------------+-------------+------------+---
      020726ae| 2986        | 343         | 50         |     <- amostras
      020726af| 2940        | 338         | 45         |

O nome do elemento fica na coluna da concentração e ocupa duas colunas
(a dele e a do erro). Célula "-" quer dizer "não detectado" e some do
gráfico, do mesmo jeito que um elemento que não aparece no .txt do XRF.

O que sai daqui tem a MESMA forma do que sai de `leitura.py` — uma lista
de amostras, cada uma com uma lista de {"z", "symbol", "valor"} —, e por
isso o resto do programa (classificação, gráficos, tabelas) não muda de
acordo com a origem dos dados. Só o "valor" muda de significado: aqui
ele é a concentração, e não a área do pico.
"""

import os
import re

from .tabela_periodica import PERIODIC_TABLE

NUMERO_ATOMICO = {simbolo: z for z, simbolo in PERIODIC_TABLE.items()}

# Como a unidade aparece no cabeçalho: "Ca (mg/kg)".
_UNIDADE = re.compile(r"\(([^)]+)\)")
# O símbolo é a primeira palavra do cabeçalho: "Ca (mg/kg)" -> "Ca".
_SIMBOLO = re.compile(r"([A-Za-z]{1,2})\b")

UNIDADE_PADRAO = "mg/kg"
ABA = "Resultados"


def _abrir(caminho):
    """A grade de células da aba "Resultados", sem as linhas vazias."""
    try:
        import openpyxl
    except ImportError:
        raise ValueError(
            "Para ler planilhas falta a biblioteca openpyxl.\n\n"
            "Instale com:  pip install openpyxl")

    if os.path.splitext(caminho)[1].lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Escolha o arquivo .xlsx exportado pelo programa "
                         "de concentrações.")

    livro = openpyxl.load_workbook(caminho, data_only=True)
    nome = next((n for n in livro.sheetnames if n.strip().lower() == ABA.lower()), None)
    if nome is None:
        raise ValueError(
            'A planilha não tem uma aba "%s" (abas encontradas: %s).'
            % (ABA, ", ".join(livro.sheetnames)))

    grade = [list(linha) for linha in livro[nome].iter_rows(values_only=True)]
    return [linha for linha in grade
            if any(c is not None and str(c).strip() for c in linha)]


def _texto(celula):
    return "" if celula is None else str(celula).strip()


def _numero(celula):
    """A célula como número. None quando está vazia, é "-" ou não é
    número — os três casos querem dizer a mesma coisa aqui: esse
    elemento não entra no gráfico desta amostra."""
    if celula is None or isinstance(celula, bool):
        return None
    if isinstance(celula, (int, float)):
        return float(celula)
    texto = str(celula).strip().replace(",", ".")
    if not texto or texto == "-":
        return None
    try:
        return float(texto)
    except ValueError:
        return None


def _achar_cabecalho(grade):
    """Onde estão os nomes dos elementos e onde começam as amostras.

    Aceita a planilha com ou sem a linha "LD": o que identifica o
    cabeçalho é a linha de "C"/"Erro", e os nomes dos elementos estão
    sempre na primeira linha do bloco.
    """
    for i in range(min(6, len(grade))):
        primeira = _texto(grade[i][0]).upper()
        if primeira == "LD" and i >= 1:
            return grade[i - 1], i + 2
        resto = [_texto(c).lower() for c in grade[i][1:] if _texto(c)]
        if resto and "erro" in resto and all(c in ("c", "erro") for c in resto):
            if i >= 1:
                return grade[i - 1], i + 1
    raise ValueError(
        'Não reconheci o cabeçalho da aba "%s". Ele deve ter o nome do '
        'elemento numa linha e "C"/"Erro" logo abaixo.' % ABA)


def _colunas_dos_elementos(linha):
    """[(coluna, símbolo)] para cada elemento do cabeçalho.

    O nome ocupa duas colunas — a da concentração e a do erro —, então
    andamos de dois em dois a partir de cada nome encontrado.
    """
    colunas = []
    j, n = 1, len(linha)
    while j < n:
        nome = _texto(linha[j])
        if not nome:
            j += 1
            continue
        achado = _SIMBOLO.match(nome)
        simbolo = achado.group(1) if achado else ""
        simbolo = simbolo[:1].upper() + simbolo[1:].lower()
        if simbolo in NUMERO_ATOMICO:
            colunas.append((j, simbolo))
        j += 2       # pula a coluna do erro
    return colunas


def _unidade(linha, colunas):
    """A unidade que está entre parênteses no cabeçalho ("mg/kg")."""
    for j, _ in colunas:
        achado = _UNIDADE.search(_texto(linha[j]))
        if achado:
            return achado.group(1).strip()
    return UNIDADE_PADRAO


def parse_planilha(caminho):
    """Lê a planilha e devolve (amostras, unidade, ignoradas).

    Cada amostra é {"code": str, "elements": [...]}, do mesmo jeito que
    o catalogador monta a partir dos .txt — o "code" é o que está na
    primeira coluna (o código do arquivo), que é justamente o que o
    arquivo de mapeamento traduz para o nome real.

    Amostra sem nenhum valor numérico fica de fora: um gráfico dela
    seria uma pizza vazia.
    """
    grade = _abrir(caminho)
    if not grade:
        raise ValueError('A aba "%s" está vazia.' % ABA)

    linha_elementos, primeira_amostra = _achar_cabecalho(grade)
    colunas = _colunas_dos_elementos(linha_elementos)
    if not colunas:
        raise ValueError(
            'Não achei nenhum elemento no cabeçalho da aba "%s". Os nomes '
            'devem ser símbolos químicos, como "Ca (mg/kg)".' % ABA)

    amostras, vazias = [], []
    for linha in grade[primeira_amostra:]:
        codigo = _texto(linha[0])
        if not codigo or codigo.upper() == "LD":
            continue
        elementos = []
        for j, simbolo in colunas:
            valor = _numero(linha[j]) if j < len(linha) else None
            if valor is None or valor <= 0:
                continue
            elementos.append({"z": NUMERO_ATOMICO[simbolo], "symbol": simbolo,
                              "valor": valor})
        if elementos:
            amostras.append({"code": codigo, "elements": elementos})
        else:
            vazias.append(codigo)

    if not amostras:
        raise ValueError(
            'Nenhuma amostra com valor numérico na aba "%s".' % ABA)

    return amostras, _unidade(linha_elementos, colunas), vazias
