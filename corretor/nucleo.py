"""Correção de resultados de FRX medidos por 3 tubos de raios X (Ag, Au, Rh).

Contexto: Ag, Au e Rh aqui são os TUBOS usados na medição (o material do
ânodo), não elementos da amostra — cada tubo produz seu próprio arquivo
de resultados, com as concentrações dos elementos reais da amostra (Ca,
Cl, Fe, K, Mn, P, Zn etc.) que aquele tubo consegue excitar/detectar. Como
tubos diferentes têm sensibilidades diferentes, o mesmo elemento medido
por 3 tubos dá 3 leituras que precisam ser reconciliadas.

Lógica (conforme especificado pelo laboratório), aplicada A CADA ELEMENTO
que aparece em pelo menos um dos 3 arquivos de tubo:

1. O tubo Au é considerado o mais confiável e serve de referência fixa.
2. Para os tubos Ag e Rh, calcula-se um "fator de correção" a partir da
   razão Au/tubo em todas as amostras da batelada (média entre a média e
   a mediana das razões) e multiplica-se Conc e Erro daquele tubo por
   esse fator.
3. Com as três leituras (Au fixo, Ag e Rh corrigidos) já na mesma escala,
   compara-se cada amostra:
   - com as 3 leituras presentes -> Teste 1 (valor médio ponderado pelo
     inverso do erro ao quadrado; descarta-se a única leitura fora de
     2 erros do valor médio; se 2 ou mais estiverem fora, descarta-se o
     maior valor entre Ag e Rh, nunca o Au);
   - com apenas 2 leituras presentes -> Teste 2 (descarta-se a leitura de
     menor prioridade, na ordem Au > Ag > Rh, quando a diferença entre
     as duas ultrapassa 2 erros combinados).
4. A tabela final tem uma linha por amostra e, para cada elemento
   encontrado, um bloco de colunas com a leitura corrigida (ou em
   branco, quando descartada) de cada tubo que o mediu — deixando claro
   de qual tubo veio cada resultado.

Cada arquivo de tubo é sempre a aba "Resultados" da planilha exportada
pelo programa de concentrações (cabeçalho em 3 linhas: nome do elemento
ocupando 2 colunas, linha "LD", depois "C"/"Erro").
"""

from __future__ import annotations

import csv
import math
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "Este programa precisa do pacote 'openpyxl' (pip install openpyxl)."
    ) from exc

from catalogador.nucleo.tabela_periodica import PERIODIC_TABLE

SIMBOLOS_VALIDOS = set(PERIODIC_TABLE.values())
_NUMERO_ATOMICO = {simbolo: z for z, simbolo in PERIODIC_TABLE.items()}

TUBOS = ("Au", "Ag", "Rh")
PRIORIDADE = {"Au": 0, "Ag": 1, "Rh": 2}  # menor = mais confiável

_PADRAO_CONC = re.compile(r"conc[a-zçã\.]*\s*\(?\s*([a-zA-Z]{1,2})\s*\)?", re.IGNORECASE)
_PADRAO_ERRO = re.compile(r"erro[a-zçã\.]*\s*\(?\s*([a-zA-Z]{1,2})\s*\)?", re.IGNORECASE)


def _normalizar_simbolo(bruto: str) -> str:
    bruto = bruto.strip()
    if len(bruto) == 1:
        return bruto.upper()
    return bruto[0].upper() + bruto[1:].lower()


def _celula(v) -> str:
    return "" if v is None else str(v).strip()


def detectar_pares(cabecalho: list) -> dict:
    """Acha, no cabeçalho, os pares (Conc, Erro) de cada elemento reconhecido.

    Retorna {simbolo: (indice_conc, indice_erro)}.
    """
    pares = {}
    n = len(cabecalho)
    for j, titulo in enumerate(cabecalho):
        if not titulo:
            continue
        m = _PADRAO_CONC.match(str(titulo).strip())
        if not m:
            continue
        simbolo = _normalizar_simbolo(m.group(1))
        if simbolo not in SIMBOLOS_VALIDOS:
            continue
        idx_erro = None
        for k in range(j + 1, min(j + 5, n)):
            titulo_k = cabecalho[k]
            if not titulo_k:
                continue
            m2 = _PADRAO_ERRO.match(str(titulo_k).strip())
            if m2 and _normalizar_simbolo(m2.group(1)) == simbolo:
                idx_erro = k
                break
        if idx_erro is not None:
            pares[simbolo] = (j, idx_erro)
    return pares


def numero(valor):
    """Converte a célula para float, aceitando vírgula decimal. None se vazia/inválida."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto or texto == "-":
        return None
    texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def casas_decimais(valor):
    """Quantas casas decimais o valor tinha originalmente (como veio do
    arquivo). None quando não dá pra saber (célula vazia)."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return 0
    if isinstance(valor, float):
        texto = repr(valor)
    else:
        texto = str(valor).strip().replace(",", ".")
    if not texto:
        return None
    if "e" in texto.lower():
        return None
    if "." in texto:
        return len(texto.split(".")[-1])
    try:
        float(texto)
    except ValueError:
        return None
    return 0


def calcular_fator(linhas: list, pares: dict, fixo: str, alvo: str):
    """Fator de correção do `alvo` em relação ao `fixo`: média(média, mediana) das razões."""
    ic, _ = pares[fixo]
    ja, _ = pares[alvo]
    razoes = []
    for linha in linhas:
        cf = numero(linha[ic]) if ic < len(linha) else None
        ca = numero(linha[ja]) if ja < len(linha) else None
        if cf is None or ca is None or ca == 0:
            continue
        razoes.append(cf / ca)
    if not razoes:
        return None
    return (statistics.fmean(razoes) + statistics.median(razoes)) / 2


def valor_medio_ponderado(dados: dict) -> float:
    """dados: {simbolo: (conc, erro)} -> média ponderada pelo inverso do erro ao quadrado."""
    numerador = sum(c / (e ** 2) for c, e in dados.values())
    denominador = sum(1 / (e ** 2) for c, e in dados.values())
    return numerador / denominador


def teste1(dados: dict) -> set:
    """Recebe as 3 leituras {tubo: (conc, erro)} e devolve o(s) tubo(s) a descartar."""
    vm = valor_medio_ponderado(dados)
    distancias = {s: abs(c - vm) - 2 * e for s, (c, e) in dados.items()}
    positivos = [s for s, d in distancias.items() if d > 0]
    if not positivos:
        return set()
    if len(positivos) == 1:
        return set(positivos)
    # 2 ou mais fora da faixa: só Ag/Rh entram na disputa, o Au nunca é descartado.
    candidatos = [s for s in ("Ag", "Rh") if s in dados]
    if not candidatos:
        return set(positivos)
    maior = max(candidatos, key=lambda s: dados[s][0])
    return {maior}


def teste2(dados: dict) -> set:
    """Recebe 2 leituras {tubo: (conc, erro)} e devolve o tubo a descartar (ou vazio)."""
    if len(dados) != 2:
        raise ValueError("teste2 espera exatamente 2 tubos")
    (s1, (c1, e1)), (s2, (c2, e2)) = list(dados.items())
    distancia = abs(c1 - c2) - 2 * math.sqrt(e1 ** 2 + e2 ** 2)
    if distancia <= 0:
        return set()
    pior = max(dados.keys(), key=lambda s: PRIORIDADE.get(s, 99))
    return {pior}


@dataclass
class Relatorio:
    fatores: dict = field(default_factory=dict)
    controle_pos_fator: dict = field(default_factory=dict)
    descartes: list = field(default_factory=list)  # [(amostra, tubo, teste)]
    tubos_presentes: tuple = ()

    def linhas_texto(self) -> list:
        linhas = [f"Tubos com dados: {', '.join(self.tubos_presentes) or '(nenhum)'}"]
        for tubo, fator in self.fatores.items():
            linhas.append(f"Fator de correção aplicado ao tubo {tubo}: {fator:.6g}")
        for tubo, (media_c, mediana_c) in self.controle_pos_fator.items():
            linhas.append(
                f"Controle pós-correção Au/{tubo} — média: {media_c:.4g}, "
                f"mediana: {mediana_c:.4g} (deveriam estar perto de 1)"
            )
        if self.descartes:
            linhas.append(f"{len(self.descartes)} leitura(s) descartada(s):")
            for amostra, tubo, teste in self.descartes:
                linhas.append(f"  - {amostra}: tubo {tubo} descartado ({teste})")
        else:
            linhas.append("Nenhuma leitura descartada.")
        return linhas


def processar_tabela(cabecalho: list, linhas: list):
    """Aplica a pipeline completa a UM elemento (cabecalho com colunas
    Conc/Erro para os tubos Ag/Au/Rh). Devolve (cabecalho, linhas_corrigidas, Relatorio)."""
    pares = detectar_pares(cabecalho)
    tubos_presentes = tuple(s for s in TUBOS if s in pares)

    largura = len(cabecalho)
    saida = [list(linha) + [None] * (largura - len(linha)) for linha in linhas]

    relatorio = Relatorio(tubos_presentes=tubos_presentes)

    fatores = {}
    if "Au" in pares:
        for tubo in ("Ag", "Rh"):
            if tubo in pares:
                fator = calcular_fator(linhas, pares, "Au", tubo)
                if fator is not None:
                    fatores[tubo] = fator
    relatorio.fatores = fatores

    for tubo, fator in fatores.items():
        ic, ie = pares[tubo]
        for linha in saida:
            c = numero(linha[ic])
            if c is None:
                continue
            e = numero(linha[ie])
            linha[ic] = c * fator
            if e is not None:
                linha[ie] = e * fator

    # Controle pós-correção (só para conferência, não realimenta o cálculo).
    for tubo in fatores:
        ic_fixo, _ = pares["Au"]
        ic_tubo, _ = pares[tubo]
        razoes = []
        for linha in saida:
            cf = numero(linha[ic_fixo])
            ca = numero(linha[ic_tubo])
            if cf is None or ca is None or ca == 0:
                continue
            razoes.append(cf / ca)
        if razoes:
            relatorio.controle_pos_fator[tubo] = (
                statistics.fmean(razoes),
                statistics.median(razoes),
            )

    id_col = 0
    for linha in saida:
        amostra = linha[id_col] if largura else None
        dados = {}
        for tubo in tubos_presentes:
            ic, ie = pares[tubo]
            c = numero(linha[ic])
            e = numero(linha[ie])
            if c is not None and e is not None:
                dados[tubo] = (c, e)

        if len(dados) == 3:
            descartados = teste1(dados)
            teste_nome = "Teste 1"
        elif len(dados) == 2:
            descartados = teste2(dados)
            teste_nome = "Teste 2"
        else:
            descartados = set()
            teste_nome = "-"

        for tubo in descartados:
            ic, ie = pares[tubo]
            linha[ic] = None
            linha[ie] = None
            relatorio.descartes.append((amostra, tubo, teste_nome))

    # Arredonda Ag/Rh pro mesmo número de casas decimais que o Au tinha
    # originalmente naquela amostra (antes de qualquer correção) — sem
    # isso, o fator de correção deixa os valores com uma precisão maior
    # do que a medição realmente tem.
    if "Au" in pares:
        ic_au, ie_au = pares["Au"]
        for linha_bruta, linha in zip(linhas, saida):
            au_conc_bruto = linha_bruta[ic_au] if ic_au < len(linha_bruta) else None
            au_erro_bruto = linha_bruta[ie_au] if ie_au < len(linha_bruta) else None
            casas_conc = casas_decimais(au_conc_bruto)
            casas_erro = casas_decimais(au_erro_bruto)
            for tubo in ("Ag", "Rh"):
                if tubo not in pares:
                    continue
                ic, ie = pares[tubo]
                valor_conc = numero(linha[ic])
                if casas_conc is not None and valor_conc is not None:
                    arredondado = round(valor_conc, casas_conc)
                    linha[ic] = int(arredondado) if casas_conc == 0 else arredondado
                valor_erro = numero(linha[ie])
                if casas_erro is not None and valor_erro is not None:
                    arredondado = round(valor_erro, casas_erro)
                    linha[ie] = int(arredondado) if casas_erro == 0 else arredondado

    return cabecalho, saida, relatorio


def _achar_cabecalho_resultados(grade: list):
    """Procura o cabeçalho de 3 linhas da aba 'Resultados' (elemento / LD /
    C-Erro), com ou sem a linha 'LD'. Devolve (linha_dos_elementos,
    indice_primeira_linha_de_dados) ou (None, None) se não achar."""
    limite = min(6, len(grade))
    for i in range(limite):
        primeira = _celula(grade[i][0]) if grade[i] else ""
        if primeira.upper() == "LD" and i >= 1:
            return grade[i - 1], i + 2
        resto = [c.strip().lower() for c in grade[i][1:] if c is not None and str(c).strip()]
        if resto and "erro" in resto and all(c in ("c", "erro") for c in resto):
            if i >= 1:
                return grade[i - 1], i + 1
    return None, None


def normalizar_bloco(grade: list):
    """Aceita tanto o formato 'Resultados' (elemento ocupando 2 colunas, LD,
    C/Erro) quanto o formato simples de uma linha só ('Conc (X)', 'Erro (X)').
    Devolve (cabecalho, linhas) no formato canônico ['Amostra', 'Conc (X)',
    'Erro (X)', ...] — um X para cada elemento real encontrado."""
    if not grade:
        return [], []

    linha_elementos, idx_dados = _achar_cabecalho_resultados(grade)
    if linha_elementos is not None:
        cabecalho = ["Amostra"]
        colunas_origem = []
        j, n = 1, len(linha_elementos)
        while j < n:
            nome = _celula(linha_elementos[j])
            if nome:
                m = re.match(r"([A-Za-z]{1,2})\b", nome)
                simbolo = _normalizar_simbolo(m.group(1)) if m else nome
                cabecalho.append(f"Conc ({simbolo})")
                colunas_origem.append(j)
                if j + 1 < n:
                    cabecalho.append(f"Erro ({simbolo})")
                    colunas_origem.append(j + 1)
                j += 2
            else:
                j += 1

        linhas = []
        for linha in grade[idx_dados:]:
            amostra = _celula(linha[0]) if linha else ""
            if not amostra or amostra.upper() == "LD":
                continue
            nova = [amostra]
            for j in colunas_origem:
                nova.append(linha[j] if j < len(linha) else None)
            linhas.append(nova)
        return cabecalho, linhas

    # Formato simples: primeira linha já é o cabeçalho "Conc (X)/Erro (X)".
    cabecalho = grade[0]
    linhas = [linha for linha in grade[1:] if _celula(linha[0])]
    return cabecalho, linhas


def abrir_resultados(caminho):
    """Abre a planilha exportada de um tubo e lê a aba 'Resultados'.
    Devolve (cabecalho, linhas) no formato canônico ['Amostra', 'Conc (X)',
    'Erro (X)', ...]."""
    caminho = Path(caminho)
    if caminho.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Escolha o arquivo .xlsx exportado pelo programa de concentrações.")

    wb = openpyxl.load_workbook(caminho, data_only=True)
    nome_aba = next((n for n in wb.sheetnames if n.strip().lower() == "resultados"), None)
    if nome_aba is None:
        raise ValueError(
            f'O arquivo não tem uma aba "Resultados" (abas encontradas: {", ".join(wb.sheetnames)}).'
        )
    ws = wb[nome_aba]
    grade = [list(linha) for linha in ws.iter_rows(values_only=True)]
    grade = [linha for linha in grade if any(v is not None and str(v).strip() for v in linha)]
    return normalizar_bloco(grade)


def aplicar_mapeamento(linhas: list, mapeamento: dict) -> list:
    """Troca o código (coluna A) de cada linha pelo nome mapeado, quando
    existir associação para ele (case-insensitive); senão mantém o código
    como está. Usado ANTES de comparar entre tubos, já que cada tubo pode
    gerar um código de arquivo diferente para a mesma amostra física."""
    if not mapeamento:
        return linhas
    novas = []
    for linha in linhas:
        nova = list(linha)
        codigo = _celula(nova[0]) if nova else ""
        if codigo:
            nova[0] = mapeamento.get(codigo.lower(), codigo)
        novas.append(nova)
    return novas


def detectar_elementos(blocos_tubo: dict) -> list:
    """blocos_tubo: {'Ag': (cabecalho, linhas) | None, 'Au': ..., 'Rh': ...}.
    Devolve a lista de todos os elementos reais encontrados em qualquer um
    dos 3 arquivos, ordenada por número atômico."""
    elementos = set()
    for bloco in blocos_tubo.values():
        if not bloco:
            continue
        cabecalho, _ = bloco
        elementos.update(detectar_pares(cabecalho).keys())
    return sorted(elementos, key=lambda simbolo: _NUMERO_ATOMICO.get(simbolo, 9999))


def montar_tabela_do_elemento(blocos_tubo: dict, elemento: str):
    """Extrai, de cada tubo, a leitura do `elemento` e monta uma tabela
    ['Amostra', 'Conc (Ag)', 'Erro (Ag)', 'Conc (Au)', 'Erro (Au)',
    'Conc (Rh)', 'Erro (Rh)'] (só com os tubos que mediram esse elemento),
    casando as amostras pelo código já mapeado. Devolve None se nenhum
    tubo tiver esse elemento."""
    dados_por_tubo = {}
    for tubo in TUBOS:
        bloco = blocos_tubo.get(tubo)
        if not bloco:
            continue
        cabecalho, linhas = bloco
        pares = detectar_pares(cabecalho)
        if elemento not in pares:
            continue
        ic, ie = pares[elemento]
        valores = {}
        for linha in linhas:
            amostra = _celula(linha[0]) if linha else ""
            if not amostra:
                continue
            c = linha[ic] if ic < len(linha) else None
            e = linha[ie] if ie < len(linha) else None
            valores[amostra] = (c, e)
        if valores:
            dados_por_tubo[tubo] = valores

    tubos_presentes = [t for t in TUBOS if t in dados_por_tubo]
    if not tubos_presentes:
        return None

    ordem_amostras = []
    vistas = set()
    for tubo in tubos_presentes:
        for amostra in dados_por_tubo[tubo]:
            if amostra not in vistas:
                vistas.add(amostra)
                ordem_amostras.append(amostra)

    cabecalho = ["Amostra"]
    for tubo in tubos_presentes:
        cabecalho += [f"Conc ({tubo})", f"Erro ({tubo})"]

    linhas = []
    for amostra in ordem_amostras:
        linha = [amostra]
        for tubo in tubos_presentes:
            c, e = dados_por_tubo[tubo].get(amostra, (None, None))
            linha += [c, e]
        linhas.append(linha)

    return cabecalho, linhas


@dataclass
class RelatorioGeral:
    por_elemento: dict = field(default_factory=dict)  # elemento -> Relatorio
    elementos_com_1_tubo: list = field(default_factory=list)

    def linhas_texto(self) -> list:
        linhas = []
        for elemento, rel in self.por_elemento.items():
            linhas.append(f"=== {elemento} ===")
            linhas.extend(rel.linhas_texto())
            linhas.append("")
        if self.elementos_com_1_tubo:
            linhas.append(
                "Só um tubo mediu (sem correção nem teste possível): "
                + ", ".join(self.elementos_com_1_tubo)
            )
        return linhas


def processar_por_tubo(blocos_tubo: dict):
    """blocos_tubo: {'Ag': (cabecalho, linhas) | None, 'Au': ..., 'Rh': ...},
    já com o mapeamento de cada tubo aplicado ao código da amostra.

    Roda a pipeline de correção em CADA elemento encontrado e monta uma
    tabela final: uma linha por amostra e, para cada elemento, um bloco de
    colunas por tubo que o mediu (corrigido, ou em branco se descartado).
    Devolve (cabecalho, linhas, RelatorioGeral).
    """
    elementos = detectar_elementos(blocos_tubo)
    relatorio_geral = RelatorioGeral()
    resultados = {}  # elemento -> {'linhas': {amostra: linha}, 'pares': {...}, 'tubos': [...]}
    ordem_amostras = []
    vistas = set()

    for elemento in elementos:
        tabela = montar_tabela_do_elemento(blocos_tubo, elemento)
        if tabela is None:
            continue
        cab_e, linhas_e = tabela
        _, linhas_corrigidas, rel = processar_tabela(cab_e, linhas_e)
        relatorio_geral.por_elemento[elemento] = rel
        if len(rel.tubos_presentes) < 2:
            relatorio_geral.elementos_com_1_tubo.append(elemento)

        pares_e = detectar_pares(cab_e)
        tubos_presentes = [t for t in TUBOS if t in pares_e]
        resultados[elemento] = {
            "linhas": {linha[0]: linha for linha in linhas_corrigidas},
            "pares": pares_e,
            "tubos": tubos_presentes,
        }
        for linha in linhas_corrigidas:
            amostra = linha[0]
            if amostra not in vistas:
                vistas.add(amostra)
                ordem_amostras.append(amostra)

    cabecalho_final = ["Amostra"]
    for elemento in elementos:
        info = resultados.get(elemento)
        if not info:
            continue
        for tubo in info["tubos"]:
            cabecalho_final.append(f"{elemento} Conc ({tubo})")
            cabecalho_final.append(f"{elemento} Erro ({tubo})")

    linhas_final = []
    for amostra in ordem_amostras:
        linha = [amostra]
        for elemento in elementos:
            info = resultados.get(elemento)
            if not info:
                continue
            linha_e = info["linhas"].get(amostra)
            for tubo in info["tubos"]:
                ic, ie = info["pares"][tubo]
                if linha_e is not None:
                    linha.append(linha_e[ic])
                    linha.append(linha_e[ie])
                else:
                    linha.append(None)
                    linha.append(None)
        linhas_final.append(linha)

    return cabecalho_final, linhas_final, relatorio_geral


def salvar_tabela(caminho, cabecalho, linhas):
    caminho = Path(caminho)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(cabecalho)
        for linha in linhas:
            ws.append(["" if v is None else v for v in linha])
        wb.save(caminho)
        return

    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(cabecalho)
        for linha in linhas:
            escritor.writerow(["" if v is None else v for v in linha])
