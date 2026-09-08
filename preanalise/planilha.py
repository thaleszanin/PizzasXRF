# -*- coding: utf-8 -*-
"""A planilha .xlsx que a pré-análise exporta.

Uma aba só, com a batelada inteira empilhada na vertical: o resumo em
cima, falando de todas as amostras, e daí pra baixo um bloco por
amostra — o nome, a tabela dos elementos e, ao lado dela, o log daquela
amostra.

    ┌───────────────────────────────────────────┬──────────────────┐
    │ Pré-análise XRF                           │                  │
    │ Resumo do conjunto                        │                  │
    │   3 de 12 amostras precisam ser refeitas… │                  │
    ├───────────────────────────────────────────┼──────────────────┤
    │ Madeira 123 — arquivo: 081025af           │ Na amostra       │
    │ Z │ Elemento │ Área │ Erro │ Erro % │ Sit.│ Madeira 123, 2   │
    │ 24│ Cr       │  310 │  968 │ 312.3  │ ERRO│ de 14 elementos  │
    │ 26│ Fe       │ 8140 │  190 │   2.3  │ ok  │ passaram de 50%… │
    └───────────────────────────────────────────┴──────────────────┘

A linha de quem estourou o limite sai em negrito, com a fonte vermelha e
um fundo rosado — o destaque tem que sobreviver a uma impressão em
preto e branco, e negrito sobrevive.

Esta camada não sabe onde o arquivo vai parar: recebe o caminho pronto
de quem chamou. E não recalcula nada — o erro de cada elemento e o texto
dos logs vêm do núcleo, os mesmos que a tela mostra.
"""

from datetime import date

from .nucleo.avaliacao import SEM_AREA, formatar_pct
from .nucleo.relatorio import linhas_do_resumo, log_da_amostra

ABA = "Pré-análise"

# As colunas do bloco de cada amostra e o quanto cada uma ocupa.
COLUNAS = (
    ("A", "Z", 6),
    ("B", "Elemento", 12),
    ("C", "Área (cps)", 15),
    ("D", "Erro (cps)", 15),
    ("E", "Erro Percentual (%)", 20),
    ("F", "Situação", 14),
)
# A coluna vazia que separa a tabela do log, e as que o log ocupa (ele é
# uma célula só, mesclada; quatro colunas dão a largura de um parágrafo).
COLUNA_VAZIA, COLUNAS_DO_LOG = "G", ("H", "I", "J", "K")
LARGURA_DO_LOG = 22
# Quantos caracteres cabem numa linha do log e quanto uma linha da
# planilha mede, em pontos. O Excel não estica sozinho a altura de uma
# célula mesclada, então a de uma amostra curta com log comprido é
# esticada na mão — senão o texto ficaria escondido.
CARACTERES_POR_LINHA, ALTURA_DA_LINHA_PT = 88, 15

# Cores pensadas pro papel branco, não pra tela: o vermelho é o do Excel
# e o verde é escuro o bastante pra ser lido impresso.
AZUL, BRANCO = "FF2C6BD4", "FFFFFFFF"
VERMELHO, ROSA = "FFC00000", "FFFCE9E9"
VERDE = "FF1F7A3D"
CINZA, FAIXA = "FF808080", "FFEAEFF9"
BORDA = "FFD9D9D9"

# O formato dos números. Área e erro são inteiros; o erro relativo vai
# com uma casa — a segunda não decide nada.
FORMATO_VALOR, FORMATO_PCT = "#,##0", "0.0"

SITUACAO_ALTA, SITUACAO_OK = "ERRO ALTO", "ok"


def _openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise ValueError(
            "Para exportar a planilha falta a biblioteca openpyxl.\n\n"
            "Instale com:  pip install openpyxl")
    return openpyxl


def _estilos():
    """Os objetos de estilo do openpyxl, criados uma vez por exportação
    (o openpyxl compartilha o mesmo objeto entre as células)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fina = Side(style="thin", color=BORDA)
    return {
        "titulo": Font(bold=True, size=16),
        "subtitulo": Font(size=10, color=CINZA),
        "secao": Font(bold=True, size=12),
        "resumo_titulo": Font(bold=True, size=11),
        "resumo_problema": Font(bold=True, color=VERMELHO),
        "resumo_ok": Font(color=VERDE),
        "resumo_normal": Font(),
        "cabecalho": Font(bold=True, color=BRANCO),
        "amostra": Font(bold=True, size=12),
        "linha": Font(),
        "linha_alta": Font(bold=True, color=VERMELHO),
        "log_problema": Font(bold=True, color=VERMELHO),
        "log_ok": Font(color=VERDE),
        "fundo_cabecalho": PatternFill("solid", fgColor=AZUL),
        "fundo_amostra": PatternFill("solid", fgColor=FAIXA),
        "fundo_alto": PatternFill("solid", fgColor=ROSA),
        "borda": Border(left=fina, right=fina, top=fina, bottom=fina),
        "esquerda": Alignment(horizontal="left", vertical="center"),
        "centro": Alignment(horizontal="center", vertical="center"),
        "direita": Alignment(horizontal="right", vertical="center"),
        "log_alinhamento": Alignment(horizontal="left", vertical="top",
                                     wrap_text=True),
    }


def _mesclar(ws, primeira, ultima, linha, texto, fonte=None, alinhamento=None):
    """Escreve um texto que ocupa várias colunas — o resumo e o título de
    cada amostra são frases, não células de tabela."""
    ws.merge_cells("%s%d:%s%d" % (primeira, linha, ultima, linha))
    celula = ws["%s%d" % (primeira, linha)]
    celula.value = texto
    if fonte is not None:
        celula.font = fonte
    if alinhamento is not None:
        celula.alignment = alinhamento
    return celula


def _preparar_aba(ws, estilos):
    """Larguras das colunas e a grade de fundo desligada — o que dá à
    planilha a cara de relatório, e não de rascunho."""
    ws.sheet_view.showGridLines = False
    for letra, _, largura in COLUNAS:
        ws.column_dimensions[letra].width = largura
    ws.column_dimensions[COLUNA_VAZIA].width = 3
    for letra in COLUNAS_DO_LOG:
        ws.column_dimensions[letra].width = LARGURA_DO_LOG


def _escrever_resumo(ws, estilos, linha, amostras, limite, mapeamento_usado,
                     falhas):
    """O log inicial: o cabeçalho da planilha e o resumo da batelada."""
    ultima = COLUNAS_DO_LOG[-1]
    _mesclar(ws, "A", ultima, linha, "Pré-análise XRF", estilos["titulo"])
    linha += 1
    _mesclar(ws, "A", ultima, linha,
             "Gerado em %s  •  nomes das amostras: %s"
             % (date.today().strftime("%d/%m/%Y"),
                "do arquivo de mapeamento" if mapeamento_usado
                else "código do arquivo (sem mapeamento carregado)"),
             estilos["subtitulo"])
    linha += 2

    _mesclar(ws, "A", ultima, linha, "Resumo do conjunto", estilos["secao"])
    linha += 1

    fontes = {"titulo": estilos["resumo_titulo"],
              "problema": estilos["resumo_problema"],
              "ok": estilos["resumo_ok"],
              "normal": estilos["resumo_normal"]}
    for papel, texto in linhas_do_resumo(amostras, limite, falhas):
        _mesclar(ws, "A", ultima, linha, texto, fontes[papel],
                 estilos["log_alinhamento"])
        linha += 1
    return linha + 1


def _escrever_amostra(ws, estilos, linha, amostra, limite):
    """Um bloco: o nome, a tabela dos elementos e o log ao lado."""
    avaliacao = amostra["avaliacao"]
    primeira = linha

    _mesclar(ws, "A", COLUNAS[-1][0], linha,
             "%s  —  arquivo: %s" % (amostra["nome"], amostra["codigo"]),
             estilos["amostra"], estilos["esquerda"])
    for letra, _, _ in COLUNAS:
        celula = ws["%s%d" % (letra, linha)]
        celula.fill = estilos["fundo_amostra"]
        celula.border = estilos["borda"]
    linha += 1

    for letra, titulo, _ in COLUNAS:
        celula = ws["%s%d" % (letra, linha)]
        celula.value = titulo
        celula.font = estilos["cabecalho"]
        celula.fill = estilos["fundo_cabecalho"]
        celula.border = estilos["borda"]
        celula.alignment = estilos["centro"]
    linha += 1

    for elemento in avaliacao["linhas"]:
        alto = elemento["alto"]
        fonte = estilos["linha_alta"] if alto else estilos["linha"]
        # o erro relativo de um elemento sem área não é número nenhum:
        # vai como texto, pra planilha não guardar um infinito
        pct = (formatar_pct(elemento["pct"]) if elemento["pct"] == SEM_AREA
               else elemento["pct"])
        valores = (
            ("A", elemento["z"], None, estilos["centro"]),
            ("B", elemento["symbol"], None, estilos["centro"]),
            ("C", elemento["area"], FORMATO_VALOR, estilos["direita"]),
            ("D", elemento["erro"], FORMATO_VALOR, estilos["direita"]),
            ("E", pct, FORMATO_PCT, estilos["direita"]),
            ("F", SITUACAO_ALTA if alto else SITUACAO_OK, None, estilos["centro"]),
        )
        for letra, valor, formato, alinhamento in valores:
            celula = ws["%s%d" % (letra, linha)]
            celula.value = valor
            celula.font = fonte
            celula.border = estilos["borda"]
            celula.alignment = alinhamento
            if formato and isinstance(valor, (int, float)):
                celula.number_format = formato
            if alto:
                celula.fill = estilos["fundo_alto"]
        linha += 1

    ultima = linha - 1
    ws.merge_cells("%s%d:%s%d" % (COLUNAS_DO_LOG[0], primeira,
                                  COLUNAS_DO_LOG[-1], ultima))
    log = ws["%s%d" % (COLUNAS_DO_LOG[0], primeira)]
    log.value = log_da_amostra(amostra["nome"], avaliacao, limite)
    log.font = (estilos["log_problema"] if avaliacao["altos"] else estilos["log_ok"])
    log.alignment = estilos["log_alinhamento"]
    precisa = -(-len(log.value) // CARACTERES_POR_LINHA)
    disponivel = ultima - primeira + 1
    if precisa > disponivel:
        ws.row_dimensions[primeira].height = (
            ALTURA_DA_LINHA_PT * (precisa - disponivel + 1))
    for l in range(primeira, ultima + 1):
        for letra in COLUNAS_DO_LOG:
            ws["%s%d" % (letra, l)].border = estilos["borda"]

    return linha + 1


def exportar(caminho, amostras, limite, mapeamento_usado=False, falhas=()):
    """Escreve a planilha e devolve o caminho.

    `amostras` é a lista de {"nome", "codigo", "avaliacao"} — a avaliação
    é a que `nucleo.avaliar` devolveu, com o mesmo limite. `falhas` são
    os arquivos que não deram pra ler, [(arquivo, motivo)]: eles não têm
    bloco, mas aparecem no resumo, senão sumiriam sem ninguém notar.
    """
    openpyxl = _openpyxl()
    livro = openpyxl.Workbook()
    ws = livro.active
    ws.title = ABA

    estilos = _estilos()
    _preparar_aba(ws, estilos)

    linha = _escrever_resumo(ws, estilos, 1, amostras, limite,
                             mapeamento_usado, falhas)
    for amostra in amostras:
        linha = _escrever_amostra(ws, estilos, linha, amostra, limite)

    livro.save(caminho)
    return caminho
